from __future__ import annotations

import csv
import base64
import os
import threading
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import mock

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PyQt6.QtCore import QThread
from PyQt6.QtWidgets import QApplication

from ui.db_widget import DBBrowserWidget, _ExportWorker, _QueryTab, _QueryWorker
from ui.db_widget import _SchemaWorker
from ui.spreadsheet_widget import SpreadsheetIO, SpreadsheetLoadWorker, SpreadsheetModel
from plugins.encrypt_plugin import _aes_gcm_decrypt, _aes_gcm_encrypt, _derive_key, _legacy_aes_gcm_decrypt
from plugins.ftp_browser_plugin import FtpProfile, _DownloadTooLarge, _FtpPanel, _MAX_DOWNLOAD_BYTES


class SpreadsheetLoadingTests(unittest.TestCase):
    def test_single_sheet_xlsx_is_opened_once_for_names_and_data(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl non installato")

        with TemporaryDirectory() as directory:
            path = Path(directory) / "single.xlsx"
            workbook = openpyxl.Workbook()
            workbook.active.append(["name", "value"])
            workbook.active.append(["one", 1])
            workbook.save(path)

            original = openpyxl.load_workbook
            with mock.patch.object(openpyxl, "load_workbook", wraps=original) as load:
                headers, rows, error, names, current = SpreadsheetIO.load_for_open(path)

            self.assertEqual(load.call_count, 1)
            self.assertEqual(headers, ["name", "value"])
            self.assertEqual(rows, [["one", "1"]])
            self.assertIsNone(error)
            self.assertEqual(names, ["Sheet"])
            self.assertEqual(current, "Sheet")

    def test_multi_sheet_inspection_defers_materialization(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl non installato")

        with TemporaryDirectory() as directory:
            path = Path(directory) / "multi.xlsx"
            workbook = openpyxl.Workbook()
            workbook.active.title = "First"
            workbook.active.append(["a"])
            workbook.create_sheet("Second").append(["b"])
            workbook.save(path)

            headers, rows, error, names, current = SpreadsheetIO.load_for_open(path)

            self.assertEqual((headers, rows, error, current), ([], [], None, ""))
            self.assertEqual(names, ["First", "Second"])

    def test_large_dataset_warning_and_owned_model(self):
        headers = ["a", "b"]
        rows = [["1", "2"], ["3", "4"]]
        with mock.patch.object(SpreadsheetIO, "LARGE_DATASET_CELLS", 4):
            self.assertIn("Dataset molto grande", SpreadsheetIO._large_dataset_warning(headers, rows))

        model = SpreadsheetModel(headers, rows, take_ownership=True)
        self.assertIs(model._data, rows)

    def test_csv_open_is_bounded_while_rows_are_accumulated(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "large.csv"
            path.write_text("a,b\n" + "\n".join(f"{i},{i}" for i in range(10)))

            with mock.patch.object(SpreadsheetIO, "OPEN_CELL_LIMIT", 6):
                headers, rows, warning, _names, _current = SpreadsheetIO.load_for_open(path)

            self.assertEqual(headers, ["a", "b"])
            self.assertEqual(rows, [["0", "0"], ["1", "1"]])
            self.assertIn("aperto parzialmente", warning)

    def test_interrupted_worker_does_not_emit_loaded_and_is_released(self):
        entered = threading.Event()

        def cancellable_load(*args):
            should_cancel = args[-1]
            entered.set()
            while not should_cancel():
                threading.Event().wait(0.001)
            return ["a"], [["1"]], None, [], ""

        emitted = []
        worker = SpreadsheetLoadWorker(Path("unused.csv"))
        worker.loaded.connect(lambda *result: emitted.append(result))
        with mock.patch.object(SpreadsheetIO, "load_for_open", side_effect=cancellable_load):
            worker.start()
            self.assertTrue(entered.wait(5))
            worker.requestInterruption()
            self.assertTrue(worker.wait(5_000))

        self.assertEqual(emitted, [])
        self.assertNotIn(worker, SpreadsheetLoadWorker._active_workers)

    def test_silent_empty_restore_creates_placeholder_but_unload_drops_result(self):
        from plugins.spreadsheet_plugin import SpreadsheetPlugin

        class Tabs:
            def __init__(self):
                self.added = []

            def add_spreadsheet_tab(self, *args):
                self.added.append(args)

        tabs = Tabs()
        plugin = SpreadsheetPlugin()
        plugin._mw = mock.Mock(_tab_manager=tabs)
        plugin._active = True
        plugin._load_jobs = {}
        plugin._loading_paths = {Path("empty.csv")}
        worker = mock.Mock()
        worker.isInterruptionRequested.return_value = False

        with mock.patch("ui.spreadsheet_widget.SpreadsheetWidget", return_value=mock.Mock()):
            plugin._on_load_ready(
                worker, Path("empty.csv"), ",", "utf-8", True, False, True,
                [], [], "", [], "")
            self.assertEqual(len(tabs.added), 1)

            plugin._active = False
            plugin._on_load_ready(
                worker, Path("late.csv"), ",", "utf-8", True, False, False,
                ["a"], [["1"]], "", [], "")

        self.assertEqual(len(tabs.added), 1)

    def test_duplicate_csv_request_does_not_repeat_open_mode_prompt(self):
        from ui.main_window import MainWindow

        with TemporaryDirectory() as directory:
            path = Path(directory) / "duplicate.csv"
            path.write_text("a,b\n1,2")

            class Tabs:
                def find_tab_by_path(self, _path):
                    return None

            plugin = mock.Mock()
            plugin._loading_paths = set()

            def start_loading(requested_path):
                plugin._loading_paths.add(requested_path)

            plugin.open_spreadsheet.side_effect = start_loading

            window = MainWindow.__new__(MainWindow)
            window._tab_manager = Tabs()
            window._spreadsheet_plugin = plugin
            window._ask_csv_open_mode = mock.Mock(return_value="spreadsheet")

            window.open_files([path, path])

            window._ask_csv_open_mode.assert_called_once_with(path)
            plugin.open_spreadsheet.assert_called_once_with(path)


class _StreamingCursor:
    description = [("id",), ("value",)]

    def __init__(self, rows, thread_ids):
        self._rows = rows
        self._position = 0
        self.fetchmany_calls = 0
        self._thread_ids = thread_ids

    def execute(self, _sql):
        self._thread_ids.add(threading.get_ident())

    def fetchmany(self, size):
        self.fetchmany_calls += 1
        batch = self._rows[self._position:self._position + size]
        self._position += len(batch)
        return batch

    def fetchall(self):
        raise AssertionError("L'export streaming non deve usare fetchall")

    def close(self):
        self._thread_ids.add(threading.get_ident())


class _StreamingConnection:
    def __init__(self, cursor, thread_ids):
        self._cursor = cursor
        self._thread_ids = thread_ids

    def cursor(self):
        self._thread_ids.add(threading.get_ident())
        return self._cursor

    def close(self):
        self._thread_ids.add(threading.get_ident())


class DatabaseExportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def test_csv_export_fetches_batches_in_worker_thread(self):
        rows = [(index, f"value-{index}") for index in range(2505)]
        thread_ids = set()
        cursor = _StreamingCursor(rows, thread_ids)
        connection = _StreamingConnection(cursor, thread_ids)
        main_thread = threading.get_ident()

        with TemporaryDirectory() as directory:
            path = Path(directory) / "export.csv"
            opts = {
                "format": "csv", "delimiter": ";", "encoding": "utf-8",
                "header": True, "limit": 0,
            }
            worker = _ExportWorker(path, opts, {"type": "sqlite"}, sql="SELECT * FROM t")
            with mock.patch("ui.db_widget._open_connection", return_value=connection):
                worker.start()
                self.assertTrue(worker.wait(10_000))

            self.assertGreaterEqual(cursor.fetchmany_calls, 4)
            self.assertNotIn(main_thread, thread_ids)
            with open(path, encoding="utf-8", newline="") as output:
                exported = list(csv.reader(output, delimiter=";"))
            self.assertEqual(exported[0], ["id", "value"])
            self.assertEqual(len(exported), len(rows) + 1)

    def test_cancelled_streaming_export_preserves_existing_file(self):
        fetch_started = threading.Event()
        release_fetch = threading.Event()

        class SlowCursor(_StreamingCursor):
            def fetchmany(self, size):
                fetch_started.set()
                release_fetch.wait(5)
                return super().fetchmany(size)

        thread_ids = set()
        cursor = SlowCursor([(index, str(index)) for index in range(2000)], thread_ids)
        connection = _StreamingConnection(cursor, thread_ids)
        with TemporaryDirectory() as directory:
            path = Path(directory) / "cancelled.csv"
            path.write_text("existing content", encoding="utf-8")
            worker = _ExportWorker(
                path, {"format": "csv", "header": True},
                {"type": "sqlite"}, sql="SELECT * FROM t")
            with mock.patch("ui.db_widget._open_connection", return_value=connection):
                worker.start()
                self.assertTrue(fetch_started.wait(5))
                worker.requestInterruption()
                release_fetch.set()
                self.assertTrue(worker.wait(10_000))

            self.assertEqual(path.read_text(encoding="utf-8"), "existing content")
            self.assertEqual(list(path.parent.iterdir()), [path])

    def test_failed_export_preserves_existing_file(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "failed.csv"
            path.write_text("existing content", encoding="utf-8")
            worker = _ExportWorker(
                path, {"format": "unsupported"}, {"type": "sqlite"},
                headers=["id"], rows=[[1]])

            worker.start()
            self.assertTrue(worker.wait(10_000))

            self.assertEqual(path.read_text(encoding="utf-8"), "existing content")
            self.assertEqual(list(path.parent.iterdir()), [path])

    def test_cancelled_export_exception_is_reported_as_cancellation(self):
        class FailingExport(_ExportWorker):
            def _write(self, path, headers, cursor):
                raise RuntimeError("driver interrupted")

            def isInterruptionRequested(self):
                return True

        with TemporaryDirectory() as directory:
            worker = FailingExport(
                Path(directory) / "cancelled.csv", {"format": "csv"},
                {"type": "sqlite"}, headers=["id"], rows=[[1]])
            completed = []
            worker.completed.connect(lambda error, count, cancelled: completed.append(
                (error, count, cancelled)))
            worker.run()

        self.assertEqual(completed, [("driver interrupted", 0, True)])


class _BlockingQueryWorker(_QueryWorker):
    def __init__(self):
        super().__init__({"type": "sqlite"}, "SELECT 1")
        self.started_event = threading.Event()

    def run(self):
        self.started_event.set()
        while not self.isInterruptionRequested():
            QThread.msleep(5)


class _BlockingExportWorker(_ExportWorker):
    def __init__(self, path):
        super().__init__(path, {"format": "csv"}, {"type": "sqlite"},
                         headers=["id"], rows=[[1]])
        self.started_event = threading.Event()

    def run(self):
        self.started_event.set()
        while not self.isInterruptionRequested():
            QThread.msleep(5)


class DatabaseWorkerLifecycleTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def _start_workers(self, tab, directory):
        query_worker = _BlockingQueryWorker()
        export_worker = _BlockingExportWorker(Path(directory) / "export.csv")
        tab._worker = query_worker
        tab._export_worker = export_worker
        query_worker.start()
        export_worker.start()
        self.assertTrue(query_worker.started_event.wait(5))
        self.assertTrue(export_worker.started_event.wait(5))
        return query_worker, export_worker

    def test_query_tab_and_browser_close_wait_for_workers(self):
        conn_info = {"type": "sqlite", "name": "test"}
        with mock.patch.object(DBBrowserWidget, "_connect"), \
             mock.patch.object(DBBrowserWidget, "_ai_on_provider_changed"):
            browser = DBBrowserWidget(conn_info, None)
        second_tab = browser._add_query_tab()

        with TemporaryDirectory() as directory:
            first_tab = browser._query_tabs.widget(0)
            first_workers = self._start_workers(first_tab, directory)
            browser._close_query_tab(0)
            self.assertTrue(all(not worker.isRunning() for worker in first_workers))
            self.assertIsNone(first_tab._worker)
            self.assertIsNone(first_tab._export_worker)

            second_workers = self._start_workers(second_tab, directory)
            browser.close()
            self.assertTrue(all(not worker.isRunning() for worker in second_workers))
            self.assertIsNone(second_tab._worker)
            self.assertIsNone(second_tab._export_worker)

        QApplication.processEvents()

    def test_export_reference_is_cleared_after_thread_finishes(self):
        tab = _QueryTab(None, {"type": "sqlite", "name": "test"}, None, 1)
        with TemporaryDirectory() as directory:
            worker = _ExportWorker(
                Path(directory) / "export.csv", {"format": "csv"},
                {"type": "sqlite"}, headers=["id"], rows=[[1]])
            tab._export_worker = worker
            worker.finished.connect(tab._on_export_thread_finished)
            worker.finished.connect(worker.deleteLater)
            worker.start()
            self.assertTrue(worker.wait(10_000))
            QApplication.processEvents()
            self.assertIsNone(tab._export_worker)
        tab.close()


class DatabaseAndPluginSecurityTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def test_schema_load_stays_in_worker_thread_and_cancels_cooperatively(self):
        thread_ids = set()

        class Connection:
            def close(self):
                thread_ids.add(threading.get_ident())

        worker = _SchemaWorker({"type": "sqlite"})
        with mock.patch("ui.db_widget._open_connection", return_value=Connection()), \
             mock.patch("ui.db_widget._get_tables", return_value=[("records", "table")]), \
             mock.patch("ui.db_widget._get_columns", return_value=[("id", "INTEGER")]):
            worker.start()
            self.assertTrue(worker.wait(5_000))
        self.assertTrue(thread_ids)
        self.assertNotIn(threading.get_ident(), thread_ids)

    def test_ftps_is_the_default_and_downloads_are_bounded(self):
        self.assertEqual(FtpProfile().protocol, "FTPS")
        panel = _FtpPanel.__new__(_FtpPanel)
        panel._worker = mock.Mock(_cancelled=threading.Event())
        ftp = mock.Mock()
        ftp.size.return_value = _MAX_DOWNLOAD_BYTES + 1
        panel._conn = ("ftps", ftp)
        with self.assertRaises(_DownloadTooLarge):
            panel._fetch_file("large.txt")
        ftp.retrbinary.assert_not_called()

    def test_ftps_uses_a_validating_context_and_private_data_channel(self):
        panel = _FtpPanel.__new__(_FtpPanel)
        ftp = mock.Mock()
        context = mock.Mock()
        with mock.patch("ftplib.FTP_TLS", return_value=ftp) as ftp_tls, \
             mock.patch("plugins.ftp_browser_plugin.ssl.create_default_context", return_value=context):
            kind, connected = panel._connect_ftp(FtpProfile(host="example.test", protocol="FTPS"))
        ftp_tls.assert_called_once_with(context=context)
        ftp.prot_p.assert_called_once_with()
        self.assertEqual((kind, connected), ("ftps", ftp))

    def test_authenticated_envelope_rejects_tampering_and_legacy_is_explicit(self):
        try:
            from cryptography.hazmat.primitives.ciphers.aead import AESGCM
        except ImportError:
            self.skipTest("cryptography non installato")
        encrypted = _aes_gcm_encrypt("confidential", "password")
        self.assertTrue(encrypted.startswith("NPQ1:AGCM:"))
        self.assertEqual(_aes_gcm_decrypt(encrypted, "password"), "confidential")
        prefix, payload = encrypted.rsplit(":", 1)
        tampered = bytearray(base64.b64decode(payload))
        tampered[-1] ^= 1
        with self.assertRaises(Exception):
            _aes_gcm_decrypt(prefix + ":" + base64.b64encode(tampered).decode("ascii"), "password")

        salt, nonce = b"s" * 16, b"n" * 12
        legacy = base64.b64encode(salt + nonce + AESGCM(_derive_key("password", salt, 32)).encrypt(
            nonce, b"old text", None)).decode("ascii")
        with self.assertRaises(ValueError):
            _aes_gcm_decrypt(legacy, "password")
        self.assertEqual(_legacy_aes_gcm_decrypt(legacy, "password"), "old text")


if __name__ == "__main__":
    unittest.main()
